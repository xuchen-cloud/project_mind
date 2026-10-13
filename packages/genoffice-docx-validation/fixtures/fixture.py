"""Controlled DOCX fixtures and independent OOXML/python-docx assertions. Never invokes LO."""
import io, json, hashlib, sys, zipfile, posixpath
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw
from lxml import etree

ROOT = Path(__file__).resolve().parent
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
OLD = '待修改中文段落：春风拂面。'
NEW = '已修改中文段落：秋风拂面。'
CELL_OLD, CELL_NEW = '待改单元', '已改单元'
assert len(OLD) == len(NEW) and len(CELL_OLD) == len(CELL_NEW)

def sha(data): return hashlib.sha256(data).hexdigest()
def save_json(name, data): (ROOT / name).write_text(json.dumps(data, ensure_ascii=False, indent=2))
def inspect(path):
    path = Path(path)
    doc = Document(path)  # Independent consumer must successfully reopen, not merely unzip.
    with zipfile.ZipFile(path) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        texts = lambda node: ''.join(node.xpath('.//w:t/text()', namespaces=NS))
        paragraphs = [texts(p) for p in xml.xpath('./w:body/w:p', namespaces=NS)]
        rels=etree.fromstring(z.read('word/_rels/document.xml.rels'))
        relmap={r.get('Id'):r.get('Target') for r in rels}
        sections = []
        section_hf=[]
        for s in xml.xpath('.//w:sectPr', namespaces=NS):
            entry = {}
            for tag in ['pgSz','pgMar','type','cols']:
                nodes = s.findall('w:' + tag, NS)
                entry[tag] = [{etree.QName(k).localname: v for k,v in n.attrib.items()} for n in nodes]
            sections.append(entry)
            refs={}
            for kind in ['header','footer']:
                for ref in s.findall('w:'+kind+'Reference', NS):
                    target=relmap[ref.get(qn('r:id'))]
                    part=posixpath.normpath(posixpath.join('word',target)) if not target.startswith('/') else target.lstrip('/')
                    refs[kind+'_'+ref.get(qn('w:type'))]=texts(etree.fromstring(z.read(part)))
            section_hf.append(refs)
        media = []
        for n in sorted(z.namelist()):
            if n.startswith('word/media/') and not n.endswith('/'):
                data=z.read(n)
                with Image.open(io.BytesIO(data)) as im:
                    media.append({'bytes_sha256':sha(data), 'size':list(im.size), 'pixels_sha256':sha(im.convert('RGBA').tobytes())})
        hf={}
        for kind in ['header','footer']:
            hf[kind]=sorted(texts(etree.fromstring(z.read(n))) for n in z.namelist() if n.startswith('word/'+kind) and n.endswith('.xml'))
        hyperlinks=sorted(r.get('Target') for r in rels if r.get('Type','').endswith('/hyperlink'))
        headings=[]
        for p in xml.xpath('./w:body/w:p', namespaces=NS):
            style=p.find('w:pPr/w:pStyle',NS)
            if style is not None and ('Heading' in style.get(qn('w:val'),'') or style.get(qn('w:val'))=='Title'):
                headings.append({'text':texts(p),'style':style.get(qn('w:val'))})
        image_extents=[dict(n.attrib) for n in xml.xpath('.//*[local-name()="extent"]')]
        ref_counts={kind:len(xml.xpath('.//w:'+kind+'Reference', namespaces=NS)) for kind in ['header','footer']}
        bold_runs=[]
        for r in xml.xpath('.//w:r',namespaces=NS):
            b=r.find('w:rPr/w:b',NS)
            if b is not None and b.get(qn('w:val')) not in ['0','false','off']:
                bold_runs.append(texts(r))
        return {'path':str(path), 'bytes':path.stat().st_size, 'sha256':sha(path.read_bytes()),
                'paragraphs':paragraphs, 'tables':[[[c.text for c in r.cells] for r in t.rows] for t in doc.tables],
                'headings':headings, 'media':media, 'image_extents':image_extents,
                'sections':sections, 'headers_footers':hf,'headers_footers_reference_counts':ref_counts,
                'headers_footers_per_section':section_hf,'explicit_bold_text':''.join(bold_runs),
                'page_break_count':len(xml.xpath('.//w:br[@w:type="page"]', namespaces=NS)),
                'hyperlinks':hyperlinks}

def generate():
    if any((ROOT/(name+'.original.docx')).exists() for name in ['simple','complex']):
        raise RuntimeError('Original fixtures already exist. Preserve them; generate only in a fresh fixture directory.')
    image=Image.new('RGB',(160,80),'#214d68'); draw=ImageDraw.Draw(image)
    draw.rectangle((12,12,70,65), fill='#e1ba6a'); draw.ellipse((90,15,140,65),fill='#f4f0e9')
    image.save(ROOT/'fixture-image.png')
    manifest={'issue':'#118/F03','edit_targets':{'paragraph':{'before':OLD,'after':NEW},'table_cell':{'table':0,'row':1,'column':1,'before':CELL_OLD,'after':CELL_NEW}},'cases':{}}
    for name in ['simple','complex']:
        doc=Document(); doc.core_properties.title='中文文档往返验证'
        doc.add_heading('中文文档往返验证',0)
        doc.add_heading('第一节 基础内容',1)
        doc.add_paragraph(OLD)
        p=doc.add_paragraph('未修改段落：山川日月保持原样。'); p.runs[0].bold=True
        table=doc.add_table(rows=2,cols=2); table.style='Table Grid'
        for row,vals in zip(table.rows,[['姓名','状态'],['张三',CELL_OLD]]):
            for cell,text in zip(row.cells,vals): cell.text=text
        doc.add_picture(str(ROOT/'fixture-image.png'), width=Inches(1.6))
        doc.add_paragraph('图片后未修改段落。')
        if name=='complex':
            section=doc.sections[0]; section.page_width=Inches(8.27); section.page_height=Inches(11.69)
            section.header.paragraphs[0].text='第一页页眉：保留中文抬头'
            section.footer.paragraphs[0].text='第一页页脚：保留中文页脚'
            doc.add_page_break(); doc.add_paragraph('分页后内容：第二页保留。')
            p=doc.add_paragraph('未修改链接：')
            rel=p.part.relate_to('https://example.com/docx-validation','http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
            link=OxmlElement('w:hyperlink'); link.set(qn('r:id'),rel)
            r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text='示例链接'; r.append(t); link.append(r); p._p.append(link)
            s=doc.add_section(WD_SECTION_START.NEW_PAGE); s.header.is_linked_to_previous=False; s.footer.is_linked_to_previous=False
            s.header.paragraphs[0].text='第二节页眉：独立抬头'; s.footer.paragraphs[0].text='第二节页脚：独立页脚'
            s.left_margin=Inches(1.25)
            doc.add_heading('第二节 分节内容',1); doc.add_paragraph('分节后内容：第三页保留。')
        path=ROOT/(name+'.original.docx'); doc.save(path)
        original=inspect(path); manifest['cases'][name]=original
    save_json('fixtures.json',manifest)
    print(json.dumps({'samples':list(manifest['cases']),'targets':manifest['edit_targets']},ensure_ascii=False))

def compare(original, candidate, edited):
    source=inspect(original); target=inspect(candidate)
    expected=dict(source)
    if edited:
        expected['paragraphs']=[p.replace(OLD,NEW) for p in source['paragraphs']]
        expected['tables']=json.loads(json.dumps(source['tables']))
        expected['tables'][0][1][1]=CELL_NEW
    checks=[]
    for key in ['paragraphs','tables','headings','image_extents','sections','headers_footers','headers_footers_reference_counts','headers_footers_per_section','explicit_bold_text','page_break_count','hyperlinks']:
        checks.append({'assertion':key,'passed':target[key]==expected[key], 'expected':expected[key],'actual':target[key]})
    def geometry(sections):
        return [{'pgSz':s['pgSz'], 'pgMar':[{k:v for k,v in n.items() if k in ['top','bottom','left','right','header','footer']} for n in s['pgMar']]} for s in sections]
    left,right=geometry(source['sections']),geometry(target['sections'])
    checks.append({'assertion':'page_size_and_main_margins','passed':left==right,'expected':left,'actual':right})
    left=[s['type'] for s in source['sections']];right=[s['type'] for s in target['sections']]
    checks.append({'assertion':'section_count_and_break_types','passed':left==right,'expected':left,'actual':right})
    for key in ['bytes_sha256','pixels_sha256','size']:
        left=sorted(str(x[key]) for x in source['media']); right=sorted(str(x[key]) for x in target['media'])
        checks.append({'assertion':'media_'+key,'passed':left==right,'expected':left,'actual':right})
    return {'status':'passed' if all(x['passed'] for x in checks) else 'failed', 'source':source,'candidate':target,'checks':checks,
            'not_verified':['visual pagination/font fidelity','GUI edit','LibreOffice external reopen (main agent owns)']}

if __name__=='__main__':
    if sys.argv[1]=='generate': generate()
    elif sys.argv[1]=='inspect': print(json.dumps(inspect(sys.argv[2]),ensure_ascii=False,indent=2))
    elif sys.argv[1]=='compare':
        result=compare(sys.argv[2],sys.argv[3],sys.argv[4]=='edited')
        Path(sys.argv[5]).write_text(json.dumps(result,ensure_ascii=False,indent=2))
        print(json.dumps({'status':result['status'],'failures':[c['assertion'] for c in result['checks'] if not c['passed']]},ensure_ascii=False))
