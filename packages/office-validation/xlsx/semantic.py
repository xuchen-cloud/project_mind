"""Additional semantic assessment; never replaces/deletes strict comparison checks."""
from pathlib import Path
from decimal import Decimal
from copy import copy
import hashlib, json, re, time
import openpyxl
from lxml import etree as ET
from compare import ROOT, xml, style, inspect

def color(c):
    return None if c is None else {'type':c.type,'value':c.value,'tint':c.tint}
def visible_style(c):
    f=c.font
    font={k:getattr(f,k) for k in ['name','sz','b','i','u','strike','vertAlign','scheme','charset','outline','shadow','condense','extend']}
    font['color']=color(f.color)
    fill={'patternType':c.fill.patternType,'fgColor':color(c.fill.fgColor)}
    if c.fill.patternType!='solid': fill['bgColor']=color(c.fill.bgColor)
    border={}
    for k in ['left','right','top','bottom','diagonal','vertical','horizontal','start','end']:
        side=getattr(c.border,k,None)
        border[k]=None if side is None or (side.style is None and side.color is None) else {'style':side.style,'color':color(side.color)}
    border.update({k:getattr(c.border,k) for k in ['diagonalUp','diagonalDown','outline']})
    return {'font':font,'fill':fill,'border':border,'alignment':xml(copy(c.alignment)),
            'protection':xml(copy(c.protection)),'number_format':c.number_format}

def cf(sheet):
    result=[]
    for obj,rules in sheet.conditional_formatting._cf_rules.items():
        for rule in rules:
            tree=rule.to_tree()
            # Only scalar numeric literals in this controlled fixture are normalized.
            for node in tree.iter():
                if ET.QName(node).localname=='formula' and re.fullmatch(r'\d+(?:\.\d+)?',node.text or ''):
                    node.text=str(Decimal(node.text).normalize())
            result.append({'range':str(obj.sqref),'rule':ET.tostring(tree,method='c14n').decode(),'dxf':xml(rule.dxf)})
    return result

def dv(sheet):
    result=[]
    for rule in sheet.data_validations.dataValidation:
        tree=rule.to_tree()
        if tree.get('errorStyle') is None: tree.set('errorStyle','stop')
        result.append(ET.tostring(tree,method='c14n').decode())
    return result

def local_tree(tree):
    return {'tag':ET.QName(tree).localname,'attrs':dict(tree.attrib),'text':tree.text or '',
            'children':[local_tree(c) for c in tree]}
def clean_ref(ref): return ref.replace('$','')
def source(book,ref):
    sheet,coords=ref.rsplit('!',1)
    sheet=sheet.strip("'").replace("''","'")
    coords=coords.replace('$','')
    return [c.value for row in book[sheet][coords] for c in row] if ':' in coords else [book[sheet][coords].value]

def chart_assessment(original,output,edited):
    b0=openpyxl.load_workbook(original,data_only=True)
    b1=openpyxl.load_workbook(output,data_only=True)
    f0=openpyxl.load_workbook(original,data_only=False)
    f1=openpyxl.load_workbook(output,data_only=False)
    reports=[]
    for name in f0.sheetnames:
        original_charts=f0[name]._charts
        actual_charts=f1[name]._charts
        for index,chart in enumerate(original_charts):
            if index>=len(actual_charts): reports.append({'sheet':name,'index':index,'status':'failed','reason':'missing chart'}); continue
            target=actual_charts[index]
            trees=[chart._write(),target._write()]
            signatures=[]
            for tree,obj in zip(trees,[chart,target]):
                refs=tree.xpath('//*[local-name()="f"]/text()')
                signatures.append({'type':type(obj).__name__, 'barDir':obj.barDir,'grouping':obj.grouping,'gapWidth':obj.gapWidth,
                    'titles':tree.xpath('//*[local-name()="title"]//*[local-name()="t"]/text()'),
                    'references':[clean_ref(r) for r in refs],
                    'axes':[local_tree(n) for n in tree.xpath('//*[local-name()="plotArea"]/*[local-name()="catAx" or local-name()="valAx"]')],
                    'anchor':local_tree(obj.anchor.to_tree()),
                    'series_order':[(s.idx,s.order) for s in obj.series]})
            checks=[{'name':k,'passed':signatures[0][k]==signatures[1][k]} for k in signatures[0]]
            cache_checks=[]
            for node in trees[1].xpath('//*[local-name()="strRef" or local-name()="numRef"]'):
                refs=node.xpath('./*[local-name()="f"]/text()')
                values=node.xpath('./*[local-name()="strCache" or local-name()="numCache"]/*[local-name()="pt"]/*[local-name()="v"]/text()')
                if refs and values:
                    expected=source(b1,refs[0])
                    actual=[float(v) for v in values] if ET.QName(node).localname=='numRef' else values
                    cache_checks.append({'reference':refs[0],'expected':expected,'actual':actual,'passed':expected==actual})
            reports.append({'sheet':name,'index':index,'edited_dependencies':edited,'status':'passed' if all(c['passed'] for c in checks+cache_checks) else 'failed',
                'properties':checks,'cache_checks':cache_checks,'signatures':{'original':signatures[0],'output':signatures[1]},
                'classification':'source references unchanged; text categories represented as strRef and caches generated from current sheet data',
                'unproven':['native application rendering and layout','pixel appearance and label overlap','behavior in Web SDK after createUnit/save','uncovered chart families and features']})
    return reports

def assess(run):
    original=ROOT/Path(run['input']['file']); output=ROOT/Path(run['output']['file'])
    b0=openpyxl.load_workbook(original,data_only=False); b1=openpyxl.load_workbook(output,data_only=False)
    checks=[]
    for name in b0.sheetnames:
        for row in b0[name]:
            for c in row:
                if c.value is None and not c.has_style: continue
                checks.append({'name':f'{name}!{c.coordinate}.effective_style','passed':visible_style(c)==visible_style(b1[name][c.coordinate])})
        checks.append({'name':f'{name}.conditional_formatting_rule_and_dxf','passed':cf(b0[name])==cf(b1[name])})
        checks.append({'name':f'{name}.data_validation','passed':dv(b0[name])==dv(b1[name])})
    charts=chart_assessment(original,output,run['branch']=='edited')
    confirmed_strict_failures=[f for f in run['failures'] if f['category'] in ['calculation','content'] and not f['name'].endswith('.charts')]
    passed=all(c['passed'] for c in checks) and all(c['status']=='passed' for c in charts) and not confirmed_strict_failures
    return {'kind':run['kind'],'family':run['family'],'branch':run['branch'],
        'strict_status_retained':run['status'],'strict_assertions_retained':run['assertions'],
        'strict_failed_assertions_retained':run['failed_assertions'],
        'supported_semantics_status':'passed' if passed else 'failed',
        'checks':checks,'charts':charts,'confirmed_content_or_calculation_failures':confirmed_strict_failures,
        'strict_difference_classification':[
          {'name':f['name'],'classification':('XML normalization/default materialization; effective tested properties preserved' if f['name'].endswith('.style') else
            'scalar numeric literal 12 to 12.0; rule, range and differential fill preserved' if f['name'].endswith('.conditional_formatting') else
            'default errorStyle stop materialized; range/list/error messages preserved' if f['name'].endswith('.data_validation') else
            'absolute reference spelling, text category reference type and caches regenerated; see independent chart checks' if f['name'].endswith('.charts') else
            'actual failed content/calculation assertion')} for f in run['failures']],
        'scope_limit':'Semantic pass covers these controlled properties only. Strict differences remain and full visual/native/WebSDK preservation is unproven.'}

if __name__=='__main__':
    start=time.perf_counter()
    strict=json.loads((ROOT/'results.json').read_text())
    results=[assess(run) for run in strict['runs'] if 'input' in run]
    default=json.loads((ROOT/'default-calculation-results.json').read_text())
    default.update(kind='simple',family='unit',branch='edited-default')
    results.append(assess(default))
    report={'strict_results_unchanged':True,'semantic_runs':results,'ms':(time.perf_counter()-start)*1000}
    (ROOT/'semantic-results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str))
    for r in results:
        print(json.dumps({k:r[k] for k in ['kind','family','branch','strict_status_retained','strict_failed_assertions_retained','supported_semantics_status']},ensure_ascii=False))
        for c in r['charts']:
            print(json.dumps({'chart_status':c['status'],'properties':c.get('properties'),'cache_checks':c.get('cache_checks')},ensure_ascii=False))
