from pathlib import Path
from zipfile import ZipFile
from lxml import etree as ET
from copy import copy
import hashlib, json, sys, time, warnings
import openpyxl

ROOT = Path(__file__).resolve().parent
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
def sha(data): return hashlib.sha256(data).hexdigest()
def xml(obj): return ET.tostring(obj.to_tree(), method='c14n').decode() if obj is not None else None
def style(cell):
    return {'font':xml(copy(cell.font)), 'fill':xml(copy(cell.fill)), 'border':xml(copy(cell.border)),
            'alignment':xml(copy(cell.alignment)), 'protection':xml(copy(cell.protection)), 'number_format':cell.number_format}
def image_info(im):
    anchor=im.anchor
    if isinstance(anchor,str): loc={'coordinate':anchor}
    else:
        loc={'type':type(anchor).__name__,'from':xml(anchor._from) if hasattr(anchor,'_from') else None,
             'to':xml(anchor.to) if hasattr(anchor,'to') else None,
             'ext':xml(anchor.ext) if hasattr(anchor,'ext') else None}
    return {'sha256':sha(im._data()),'width':im.width,'height':im.height,'anchor':loc}
def cf_info(sheet):
    return [{'range':str(cf.sqref),'rules':[xml(rule) for rule in rules]}
            for cf,rules in sheet.conditional_formatting._cf_rules.items()]
def dimensions(sheet):
    return {'rows':{k:{'height':v.height,'hidden':v.hidden} for k,v in sheet.row_dimensions.items() if v.height is not None or v.hidden},
            'cols':{k:{'width':v.width,'hidden':v.hidden} for k,v in sheet.column_dimensions.items()}}
def inspect(path):
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter('always')
        book=openpyxl.load_workbook(path,data_only=False)
        cached=openpyxl.load_workbook(path,data_only=True)
    with ZipFile(path) as z:
        parts={i.filename:sha(z.read(i.filename)) for i in z.infolist()}
        sheet_xml={n:ET.fromstring(z.read(n)) for n in parts if n.startswith('xl/worksheets/sheet') and n.endswith('.xml')}
        raw_caches={n:{c.get('r'):{'formula':c.find('s:f',NS).text, 'value':c.find('s:v',NS).text if c.find('s:v',NS) is not None else None}
                       for c in tree.findall('.//s:c',NS) if c.find('s:f',NS) is not None} for n,tree in sheet_xml.items()}
        chart_parts={n:ET.fromstring(z.read(n)) for n in parts if n.startswith('xl/charts/chart') and n.endswith('.xml')}
        chart_details={}
        for n,tree in chart_parts.items():
            def texts(xpath): return tree.xpath(xpath)
            chart_details[n]={
                'titles':texts('//*[local-name()="title"]//*[local-name()="t"]/text()'),
                'references':texts('//*[local-name()="f"]/text()'),
                'cache_values':texts('//*[local-name()="pt"]/*[local-name()="v"]/text()'),
                'chart_type':[ET.QName(e).localname for e in tree.xpath('//*[local-name()="plotArea"]/*') if ET.QName(e).localname.endswith('Chart')],
                'axes': [ET.tostring(e,method='c14n').decode() for e in tree.xpath('//*[local-name()="plotArea"]/*[local-name()="catAx" or local-name()="valAx"]')],
            }
    sheets={}
    for sheet in book:
        cells={c.coordinate:c for row in sheet for c in row if c.value is not None or c.has_style}
        sheets[sheet.title]={'state':sheet.sheet_state,'values':{a:c.value for a,c in cells.items()},
            'caches':{a:cached[sheet.title][a].value for a,c in cells.items() if c.data_type=='f'},
            'styles':{a:style(c) for a,c in cells.items()},'merges':sorted(str(x) for x in sheet.merged_cells.ranges),
            'freeze':sheet.freeze_panes,'dimensions':dimensions(sheet),'conditional_formatting':cf_info(sheet),
            'data_validation':xml(sheet.data_validations), 'cf_differential_styles':[xml(r.dxf) for rules in sheet.conditional_formatting._cf_rules.values() for r in rules],
            'charts':[xml(c) for c in sheet._charts], 'images':[image_info(im) for im in sheet._images]}
    return {'sheets':sheets,'order':book.sheetnames,'parts':parts,'raw_formula_caches':raw_caches,'chart_details':chart_details,'warnings':[str(w.message) for w in caught]}

def compare(input_path,output_path,edited):
    before,after=inspect(input_path),inspect(output_path)
    checks=[]
    def check(name,expected,actual,category='preservation'):
        checks.append({'name':name,'category':category,'passed':expected==actual,'expected':expected,'actual':actual})
    check('sheet_order',before['order'],after['order'],'content')
    for name,source in before['sheets'].items():
        if name not in after['sheets']:
            check(f'{name}.exists',True,False,'content'); continue
        target=after['sheets'][name]
        expected_values=dict(source['values'])
        if edited and name=='明细': expected_values.update({'A3':'香蕉（SDK编辑）','B3':15,'B5':'=SUM(B3:B4)+5'})
        for coord,value in expected_values.items(): check(f'{name}!{coord}.value_or_formula',value,target['values'].get(coord),'content')
        extras={a:v for a,v in target['values'].items() if a not in expected_values and v is not None}
        check(f'{name}.unexpected_nonempty_cells',{},extras,'content')
        expected_caches=dict(source['caches'])
        if edited:
            if name=='明细': expected_caches.update({'B5':40,'B7':30})
            if name=='说明': expected_caches.update({'B3':40})
        for coord,value in expected_caches.items(): check(f'{name}!{coord}.formula_cache',value,target['caches'].get(coord),'calculation')
        for key in ['state','merges','freeze','dimensions','conditional_formatting','cf_differential_styles','data_validation','charts','images']:
            check(f'{name}.{key}',source[key],target[key], 'content' if key in ['state','merges','images','charts'] else 'preservation')
        for coord,value in source['styles'].items(): check(f'{name}!{coord}.style',value,target['styles'].get(coord),'style')
    deleted=sorted(set(before['parts'])-set(after['parts']))
    added=sorted(set(after['parts'])-set(before['parts']))
    changed=sorted(n for n in set(before['parts'])&set(after['parts']) if before['parts'][n]!=after['parts'][n])
    failures=[c for c in checks if not c['passed']]
    critical=[c for c in failures if c['category'] in ['content','calculation']]
    return {'status':'failed' if critical else ('limited' if failures else 'passed'),
            'input':{'file':Path(input_path).name,'sha256':sha(input_path.read_bytes())},
            'output':{'file':Path(output_path).name,'sha256':sha(output_path.read_bytes())},
            'assertions':len(checks),'passed_assertions':len(checks)-len(failures),'failed_assertions':len(failures),
            'checks':checks,'failures':failures,'zip_parts':{'deleted':deleted,'added':added,'changed':changed},
            'raw_formula_caches':{'original':before['raw_formula_caches'],'output':after['raw_formula_caches']},
            'chart_details':{'original':before['chart_details'],'output':after['chart_details']},'warnings':after['warnings']}

if __name__=='__main__':
    start=time.perf_counter()
    sdk=json.loads((ROOT/'sdk-results.json').read_text())
    manifest=json.loads((ROOT/'fixtures.json').read_text())
    results=[]
    for run in sdk['runs']:
        result={k:run.get(k) for k in ['kind','family','branch']}
        if 'output' not in run:
            result.update({'status':'failed','error':run.get('error')})
        else:
            try: result.update(compare(ROOT/f"{run['kind']}.original.xlsx",ROOT/Path(run['output']['file']),run['branch']=='edited'))
            except Exception as e: result.update({'status':'failed','error':{'name':type(e).__name__,'message':str(e)}})
        results.append(result)
    unchanged={kind:sha((ROOT/Path(f['file'])).read_bytes())==f['sha256'] for kind,f in manifest['samples'].items()}
    report={'issue':'#118/F03','sdk':sdk,'independent_validator':'openpyxl 3.1.5 + lxml ZIP/XML',
            'original_checksums_unchanged':unchanged,'runs':results,'comparison_ms':round((time.perf_counter()-start)*1000,3)}
    (ROOT/'results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str))
    for r in results:
        print(json.dumps({k:r.get(k) for k in ['kind','family','branch','status','assertions','failed_assertions','error']},ensure_ascii=False))
        for f in r.get('failures',[]): print('  '+f['name'])
