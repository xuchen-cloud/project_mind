from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib, json, platform, time, sys
from lxml import etree as ET
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image
from PIL import Image as PILImage, ImageDraw

ROOT = Path(__file__).resolve().parent
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def seed_caches(path, values):
    # openpyxl cannot calculate formulas: known fixture results are seeded explicitly.
    with ZipFile(path) as z:
        parts = {i.filename: z.read(i.filename) for i in z.infolist()}
    for name, cells in values.items():
        root = ET.fromstring(parts[name])
        for address, result in cells.items():
            cell = root.find(f'.//s:c[@r="{address}"]', NS)
            assert cell is not None and cell.find('s:f', NS) is not None
            value = cell.find('s:v', NS)
            if value is None:
                value = ET.SubElement(cell, '{'+NS['s']+'}v')
            value.text = str(result)
        parts[name] = ET.tostring(root)
    with ZipFile(path, 'w', ZIP_DEFLATED) as z:
        for name, content in parts.items():
            z.writestr(name, content)

def generate(kind):
    book = Workbook()
    sheet = book.active
    sheet.title = '明细'
    notes = book.create_sheet('说明')
    sheet['A1'] = '中文销售统计'
    sheet.merge_cells('A1:C1')
    sheet['A2'], sheet['B2'], sheet['C2'] = '品名', '金额', '状态'
    sheet['A3'], sheet['B3'], sheet['C3'] = '苹果', 10, '已完成'
    sheet['A4'], sheet['B4'], sheet['C4'] = '梨子', 20, '待处理'
    sheet['A5'], sheet['B5'] = '合计', '=SUM(B3:B4)'
    sheet['A7'], sheet['B7'] = '未编辑公式（依赖B3）', '=B3*2'
    sheet['A8'], sheet['B8'] = '保留百分比', 0.125
    sheet['A9'], sheet['B9'] = '保留中文及换行', '原样保留\n第二行'
    for row in (3, 4, 5, 7):
        sheet.cell(row, 2).number_format = '#,##0.00'
    sheet['B8'].number_format = '0.0%'
    sheet['B9'].alignment = Alignment(wrap_text=True, vertical='center')
    sheet['A1'].font = Font(name='Arial', size=16, bold=True, color='FF17365D')
    sheet['A1'].fill = PatternFill('solid', fgColor='FFE2EFDA')
    for cell in sheet[2]:
        cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFFFF')
        cell.fill = PatternFill('solid', fgColor='FF305496')
        cell.alignment = Alignment(horizontal='center')
        cell.border = Border(bottom=Side(style='thin', color='FF000000'))
    sheet.column_dimensions['A'].width = 28
    sheet.column_dimensions['B'].width = 22
    sheet.column_dimensions['C'].width = 16
    sheet.row_dimensions[1].height = 26
    sheet.row_dimensions[9].height = 32
    sheet.freeze_panes = 'B3'
    notes['A1'] = '未编辑工作表：请保留中文'
    notes['A2'], notes['B2'] = '固定值', 123.456
    notes['B2'].number_format = '0.000'
    notes['A3'], notes['B3'] = '跨表合计', "='明细'!B5"
    caches = {'xl/worksheets/sheet1.xml': {'B5': 30, 'B7': 20},
              'xl/worksheets/sheet2.xml': {'B3': 30}}
    if kind == 'complex':
        sheet.conditional_formatting.add('B3:B4', CellIsRule(operator='greaterThan', formula=['12'], fill=PatternFill('solid', fgColor='FFFFC7CE')))
        dv = DataValidation(type='list', formula1='"已完成,待处理"', allow_blank=False)
        dv.errorTitle, dv.error, dv.showErrorMessage = '无效状态', '请选择列表中的状态', True
        sheet.add_data_validation(dv)
        dv.add('C3:C4')
        chart = BarChart()
        chart.title, chart.y_axis.title, chart.x_axis.title = '未编辑金额图', '金额', '品名'
        chart.add_data(Reference(sheet, min_col=2, min_row=2, max_row=4), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=3, max_row=4))
        sheet.add_chart(chart, 'E2')
        picture = PILImage.new('RGB', (120, 60), '#305496')
        ImageDraw.Draw(picture).rectangle((10, 10, 110, 50), fill='#E2EFDA')
        image_path = ROOT / 'fixture-image.png'
        picture.save(image_path)
        sheet.add_image(Image(image_path), 'E18')
        hidden = book.create_sheet('隐藏数据')
        hidden['A1'], hidden['B1'] = '未触碰隐藏资料', 99
        hidden.sheet_state = 'hidden'
    path = ROOT / f'{kind}.original.xlsx'
    book.save(path)
    seed_caches(path, caches)
    loaded = openpyxl.load_workbook(path, data_only=True)
    assert loaded['明细']['B5'].value == 30 and loaded['明细']['B7'].value == 20
    return {'file': path.name, 'sha256': sha(path), 'bytes': path.stat().st_size,
            'sheets': book.sheetnames, 'seeded_formula_caches': caches,
            'edits': {'明细': {'A3': '香蕉（SDK编辑）', 'B3': 15, 'B5': '=SUM(B3:B4)+5'}},
            'edited_formula_results': {'明细': {'B5': 40, 'B7': 30}, '说明': {'B3': 40}}}

if __name__ == '__main__':
    if (ROOT/'fixtures.json').exists():
        existing=json.loads((ROOT/'fixtures.json').read_text())
        assert all(sha(ROOT/Path(s['file']))==s['sha256'] for s in existing['samples'].values()), 'Original fixture checksum changed'
        print('Reusing existing immutable original fixtures; checksums verified.')
        sys.exit(0)
    if any((ROOT/f'{kind}.original.xlsx').exists() for kind in ['simple', 'complex']):
        raise RuntimeError('Original fixtures exist without metadata; refusing to overwrite.')
    start = time.perf_counter()
    manifest = {'issue': '#118/F03', 'python': platform.python_version(), 'openpyxl': openpyxl.__version__,
                'samples': {kind: generate(kind) for kind in ('simple', 'complex')}}
    manifest['generation_ms'] = round((time.perf_counter()-start)*1000, 3)
    (ROOT/'fixtures.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2))
    print(json.dumps(manifest, ensure_ascii=False, indent=2))
