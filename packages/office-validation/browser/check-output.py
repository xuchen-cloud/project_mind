"""Read existing Office bytes only. Writes verification.json and verification.md here."""
import sys
sys.dont_write_bytecode = True
from pathlib import Path
from zipfile import ZipFile
from copy import copy
from collections import Counter
from importlib.metadata import version
from datetime import datetime, timezone
import hashlib, json, time, warnings
from lxml import etree as E
import openpyxl
from docx import Document
from pptx import Presentation

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent
PREFIX='浏览器编辑：'
def sha(b): return hashlib.sha256(b).hexdigest()
def tree(node):
    if node is None: return None
    return {'tag':E.QName(node).localname,'attrs':{E.QName(k).localname:v for k,v in node.attrib.items()},
            'text':node.text or '', 'children':[tree(c) for c in node]}
def obj(o): return tree(copy(o).to_tree()) if o is not None else None
def package(path):
    with ZipFile(path) as z:
        names=z.namelist()
        assert z.testzip() is None, 'ZIP integrity failure'
        return {n:z.read(n) for n in names}
def media(parts,prefix):
    return sorted(sha(v) for k,v in parts.items() if k.startswith(prefix) and not k.endswith('/'))
def core(parts):
    root=E.fromstring(parts['docProps/core.xml']) if 'docProps/core.xml' in parts else None
    return {E.QName(c).localname:c.text or '' for c in root} if root is not None else {}
def stable_core(parts):
    return {k:v for k,v in core(parts).items() if k not in ['created','modified','revision','lastModifiedBy'] and v}
def app_metadata(parts):
    node=E.fromstring(parts['docProps/app.xml']) if 'docProps/app.xml' in parts else None
    return {E.QName(c).localname:c.text or '' for c in node if E.QName(c).localname in ['Company','Manager','HyperlinkBase'] and c.text} if node is not None else {}
def color(c): return None if c is None else {'type':c.type,'value':c.value,'tint':c.tint}
def cell_style(c):
    font={k:getattr(c.font,k) for k in ['name','sz','b','i','u','strike','vertAlign','scheme','charset','outline','shadow']}
    font['color']=color(c.font.color)
    fill={'patternType':c.fill.patternType,'fgColor':color(c.fill.fgColor)}
    if c.fill.patternType!='solid': fill['bgColor']=color(c.fill.bgColor)
    border={k:(None if getattr(c.border,k,None) is None or (getattr(c.border,k).style is None and getattr(c.border,k).color is None)
             else {'style':getattr(c.border,k).style,'color':color(getattr(c.border,k).color)}) for k in ['left','right','top','bottom','diagonal']}
    return {'font':font,'fill':fill,'border':border,'alignment':obj(c.alignment),'protection':obj(c.protection),'number_format':c.number_format}

class Checks:
    def __init__(self): self.rows=[]
    def eq(self,name,expected,actual,category='preservation'):
        self.rows.append({'name':name,'category':category,'status':'passed' if expected==actual else 'failed','expected':expected,'actual':actual})
    def diagnostic(self,name,expected,actual,category='serialization'):
        self.rows.append({'name':name,'category':category,'status':'unchanged' if expected==actual else 'limited','expected':expected,'actual':actual})
    def known(self,name,original,baseline,actual):
        self.rows.append({'name':name,'category':'known_conversion_limit','status':'known_limit' if baseline==actual else 'failed',
                          'original':original,'node_noop_baseline':baseline,'actual':actual})

def xlsx(source,target,c):
    a=openpyxl.load_workbook(source,data_only=False); b=openpyxl.load_workbook(target,data_only=False)
    av=openpyxl.load_workbook(source,data_only=True); bv=openpyxl.load_workbook(target,data_only=True)
    c.eq('sheet_order',a.sheetnames,b.sheetnames)
    for s in a:
        if s.title not in b.sheetnames: c.eq(s.title+'.exists',True,False); continue
        t=b[s.title]
        values={v.coordinate:v.value for row in s for v in row if v.value is not None}
        expected=dict(values)
        if s.title=='明细': expected.update(A1='浏览器编辑',A3='键盘中文编辑')
        actual={v.coordinate:v.value for row in t for v in row if v.value is not None}
        c.eq(s.title+'.values_and_formula_strings',expected,actual,'content')
        for address in values:
            original=s[address]; actualcell=t[address]
            c.eq(f'{s.title}!{address}.effective_style',cell_style(original),cell_style(actualcell),'style')
            c.diagnostic(f'{s.title}!{address}.strict_style',
                         {'font':obj(original.font),'fill':obj(original.fill),'border':obj(original.border)},
                         {'font':obj(actualcell.font),'fill':obj(actualcell.fill),'border':obj(actualcell.border)})
            if original.data_type=='f':
                c.eq(f'{s.title}!{address}.formula_cache',av[s.title][address].value,bv[s.title][address].value,'calculation')
        for name,val,val2 in [('visibility',s.sheet_state,t.sheet_state),('merged',sorted(map(str,s.merged_cells.ranges)),sorted(map(str,t.merged_cells.ranges))),('freeze',s.freeze_panes,t.freeze_panes)]:
            c.eq(s.title+'.'+name,val,val2,'layout')
        def dims(sheet):
            return {'rows':{str(k):{'height':v.height,'hidden':v.hidden} for k,v in sheet.row_dimensions.items() if v.height is not None or v.hidden},
                    'cols':{k:{'width':v.width,'hidden':v.hidden} for k,v in sheet.column_dimensions.items()}}
        c.eq(s.title+'.dimensions',dims(s),dims(t),'layout')
        c.eq(s.title+'.objects',{'charts':len(s._charts),'images':len(s._images),'validation':obj(s.data_validations),'conditional_formatting':len(s.conditional_formatting)},
                               {'charts':len(t._charts),'images':len(t._images),'validation':obj(t.data_validations),'conditional_formatting':len(t.conditional_formatting)})
    def xml_caches(p):
        result={}
        for name,data in p.items():
            if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
                node=E.fromstring(data)
                result[name]={x.get('r'):{'formula':x.xpath('./*[local-name()="f"]/text()'),'cache':x.xpath('./*[local-name()="v"]/text()')}
                              for x in node.xpath('//*[local-name()="c"]') if x.xpath('./*[local-name()="f"]')}
        return result
    c.eq('raw_XML_formula_cache',xml_caches(package(source)),xml_caches(package(target)),'calculation')
    c.eq('known_formula_values',{'明细!B5':30,'明细!B7':20,'说明!B3':30},
         {'明细!B5':bv['明细']['B5'].value,'明细!B7':bv['明细']['B7'].value,'说明!B3':bv['说明']['B3'].value},'calculation')

def docx(source,target,c):
    a=Document(source); b=Document(target)
    texts=[p.text for p in a.paragraphs]; expected=list(texts); expected[0]=PREFIX+expected[0]
    c.eq('paragraph_text_order_and_prefix',expected,[p.text for p in b.paragraphs],'content')
    def tables(doc):
        return [{'rows':len(t.rows),'columns':len(t.columns),'cells':[[cell.text for cell in row.cells] for row in t.rows],
                 'widths':[[cell.width for cell in row.cells] for row in t.rows]} for t in doc.tables]
    c.eq('tables_content_geometry',tables(a),tables(b),'content')
    def blocks(doc):
        return [{'kind':E.QName(n).localname,'paragraphs':[''.join(p.xpath('.//*[local-name()="t"]/text()')) for p in n.xpath('.//*[local-name()="p"]')]
                 if E.QName(n).localname=='tbl' else [''.join(n.xpath('.//*[local-name()="t"]/text()'))]} for n in doc.element.body if E.QName(n).localname in ['p','tbl']]
    expected_blocks=blocks(a); expected_blocks[0]['paragraphs'][0]=PREFIX+expected_blocks[0]['paragraphs'][0]
    c.eq('full_body_block_order',expected_blocks,blocks(b),'content')
    def sections(doc):
        return [{k:int(getattr(s,k)) if getattr(s,k) is not None else None for k in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin','header_distance','footer_distance','gutter','orientation','start_type']} for s in doc.sections]
    c.eq('page_section_parameters',sections(a),sections(b),'layout')
    def image_layout(doc):
        return [{'width':int(x.width),'height':int(x.height)} for x in doc.inline_shapes]
    c.eq('image_extents',image_layout(a),image_layout(b),'layout')
    c.eq('paragraph_style_names',[p.style.name for p in a.paragraphs],[p.style.name for p in b.paragraphs],'style')
    def headers(doc): return [{'header':[p.text for p in s.header.paragraphs],'footer':[p.text for p in s.footer.paragraphs]} for s in doc.sections]
    c.eq('headers_footers',headers(a),headers(b),'content')
    pa=package(source); pb=package(target)
    c.eq('images_bytes',media(pa,'word/media/'),media(pb,'word/media/'),'content')
    def links(p):
        result=[]
        for name,data in p.items():
            if name.endswith('.rels') and name.startswith('word/'):
                result.extend(sorted(x.get('Target') for x in E.fromstring(data) if x.get('Type','').endswith('/hyperlink')))
        return sorted(result)
    c.eq('hyperlinks',links(pa),links(pb),'content')
    def page_breaks(p): return len(E.fromstring(p['word/document.xml']).xpath('//*[local-name()="br" and @*[local-name()="type"]="page"]'))
    c.eq('explicit_page_break_count',page_breaks(pa),page_breaks(pb),'layout')

def pptx(source,target,c):
    a=Presentation(source); b=Presentation(target)
    c.eq('slide_dimensions',[a.slide_width,a.slide_height],[b.slide_width,b.slide_height],'layout')
    c.eq('slide_count',len(a.slides),len(b.slides),'content')
    for index,s in enumerate(a.slides):
        if index>=len(b.slides): continue
        t=b.slides[index]
        c.eq(f'slide{index}.shape_count',len(s.shapes),len(t.shapes),'content')
        for k,x in enumerate(s.shapes):
            if k>=len(t.shapes): continue
            y=t.shapes[k]; base=f'slide{index}.shape{k}'
            c.eq(base+'.name_type',[x.name,int(x.shape_type)],[y.name,int(y.shape_type)],'metadata')
            c.eq(base+'.geometry_rotation',[x.left,x.top,x.width,x.height,x.rotation],[y.left,y.top,y.width,y.height,y.rotation],'layout')
            c.eq(base+'.text_frame_presence',x.has_text_frame,y.has_text_frame,'content')
            if x.has_text_frame and y.has_text_frame:
                expected=x.text
                if index==0 and k==0:
                    first=x.text_frame.paragraphs[0].runs[0].text
                    expected='浏览器编辑'+expected[len(first):]
                c.eq(base+'.text',expected,y.text,'edit' if index==0 and k==0 else 'content')
                def runs(sh):
                    return [[{'text':r.text,'font_size':r.font.size,'font_name':r.font.name,'bold':r.font.bold,'italic':r.font.italic,'underline':r.font.underline}
                             for r in p.runs] for p in sh.text_frame.paragraphs]
                r0=runs(x); r1=runs(y)
                if index==0 and k==0: r0[0][0]['text']='浏览器编辑'
                c.eq(base+'.run_text_font_attributes',r0,r1,'style')
                c.eq(base+'.paragraph_alignment',[p.alignment for p in x.text_frame.paragraphs],[p.alignment for p in y.text_frame.paragraphs],'layout')
            if x.shape_type==13 and y.shape_type==13:
                c.eq(base+'.picture_bytes',sha(x.image.blob),sha(y.image.blob),'content')
                c.eq(base+'.picture_crop',[x.crop_left,x.crop_right,x.crop_top,x.crop_bottom],[y.crop_left,y.crop_right,y.crop_top,y.crop_bottom],'layout')
            if x.has_table:
                c.eq(base+'.table',[[z.text for z in row.cells] for row in x.table.rows],[[z.text for z in row.cells] for row in y.table.rows] if y.has_table else None,'content')
        c.eq(f'slide{index}.notes',s.notes_slide.notes_text_frame.text if s.has_notes_slide else None,t.notes_slide.notes_text_frame.text if t.has_notes_slide else None,'content')
    pa=package(source); pb=package(target)
    c.eq('media_bytes',media(pa,'ppt/media/'),media(pb,'ppt/media/'),'content')
    def masters(parts):
        return {name:tree(E.fromstring(data)) for name,data in parts.items() if name.startswith(('ppt/slideMasters/','ppt/slideLayouts/')) and name.endswith('.xml')}
    baseline=ROOT/'pptx/simple.product-unchanged.pptx'
    if baseline.exists():
        c.known('master_layout_known_conversion_loss',masters(pa),masters(package(baseline)),masters(pb))
    else: c.diagnostic('master_layout_original_difference',masters(pa),masters(pb),'known_conversion_limit')
    def slide_metadata(parts):
        node=E.fromstring(parts['ppt/presentation.xml'])
        return {k:[tree(x) for x in node.xpath(f'./*[local-name()="{k}"]')] for k in ['sldSz','notesSz','defaultTextStyle']}
    c.eq('presentation_page_metadata',slide_metadata(pa),slide_metadata(pb),'metadata')

def verify(engine,format):
    candidate=HERE/'output'/engine/f'{format}.browser-edited.{format}'
    source=ROOT/format/f'simple.original.{format}'
    result={'engine':engine,'format':format,'source':str(source),'candidate':str(candidate)}
    if not candidate.exists(): return {**result,'status':'pending','unlock_condition':'Generate candidate at the specified path, then rerun this script.'}
    start=time.perf_counter(); c=Checks()
    try:
        before=package(source); after=package(candidate)
        result.update({'original_sha256':sha(source.read_bytes()),'candidate_sha256':sha(candidate.read_bytes()),'bytes':candidate.stat().st_size})
        with warnings.catch_warnings(record=True) as warns:
            warnings.simplefilter('always')
            {'xlsx':xlsx,'docx':docx,'pptx':pptx}[format](source,candidate,c)
        result['library_warnings']=[str(w.message) for w in warns]
        c.eq('stable_core_metadata',stable_core(before),stable_core(after),'metadata')
        c.eq('authored_application_metadata',app_metadata(before),app_metadata(after),'metadata')
        c.diagnostic('core_metadata_full_including_timestamps',core(before),core(after),'metadata')
        baselines={'xlsx':ROOT/'xlsx/simple.unit.noop.xlsx','docx':ROOT/'docx/simple.traditional.unit-noop.docx','pptx':ROOT/'pptx/simple.product-unchanged.pptx'}
        baseline=baselines[format]
        result['node_noop_baseline']=str(baseline)
        baseline_checks=Checks()
        if baseline.exists():
            with warnings.catch_warnings():
                warnings.simplefilter('ignore')
                {'xlsx':xlsx,'docx':docx,'pptx':pptx}[format](source,baseline,baseline_checks)
            bp=package(baseline)
            baseline_checks.eq('stable_core_metadata',stable_core(before),stable_core(bp),'metadata')
            baseline_checks.eq('authored_application_metadata',app_metadata(before),app_metadata(bp),'metadata')
            by_name={x['name']:x for x in baseline_checks.rows}
            result['node_noop_baseline_sha256']=sha(baseline.read_bytes())
            for check in c.rows:
                if check['status']=='failed' and check['name'] in by_name:
                    actual0=by_name[check['name']]['actual']
                    check['node_noop_actual']=actual0
                    check['provenance']='preexisting_node_conversion_difference' if actual0==check['actual'] else 'additional_difference_after_WebSDK_save'
                    if check['name']=='page_section_parameters':
                        expected=[{**v,'gutter':v['gutter'] or 0} for v in check['expected']]
                        actual=[{**v,'gutter':v['gutter'] or 0} for v in check['actual']]
                        check['independent_semantic_assessment']='equivalent_absent_gutter_is_zero' if expected==actual else 'actual_page_parameter_difference'
                    if check['name'].endswith('.paragraph_alignment') and format=='pptx':
                        default=E.fromstring(before['ppt/presentation.xml']).xpath('//*[local-name()="defaultTextStyle"]/*[local-name()="lvl1pPr"]/@algn')
                        if default==['l'] and check['expected']==[None] and check['actual']==[1]:
                            check['independent_semantic_assessment']='source_level_1_default_left_materialized_as_explicit_left'
            result['additional_failed_fields_vs_node_noop']=[x['name'] for x in c.rows if x.get('provenance')=='additional_difference_after_WebSDK_save']
        result['edit_status']='passed' if all(x['status']=='passed' for x in c.rows if x['name'] in ['明细.values_and_formula_strings','paragraph_text_order_and_prefix','full_body_block_order','slide0.shape0.text']) else 'failed'
        result['zip_parts']={'deleted':sorted(set(before)-set(after)),'added':sorted(set(after)-set(before)),
                             'changed':sorted(n for n in set(before)&set(after) if before[n]!=after[n])}
        result['status']='failed' if any(x['status']=='failed' for x in c.rows) else ('limited' if any(x['status'] in ['limited','known_limit'] for x in c.rows) else 'passed')
    except Exception as e:
        result.update(status='failed',error={'type':type(e).__name__,'message':str(e)})
    result.update(checks=c.rows,counts=dict(Counter(x['status'] for x in c.rows)),ms=round((time.perf_counter()-start)*1000,3))
    result['failures']=[x for x in c.rows if x['status']=='failed']
    return result

if __name__=='__main__':
    started=time.perf_counter()
    runs=[verify(engine,fmt) for engine in ['chromium','webkit'] for fmt in ['xlsx','docx','pptx']]
    report={'task':'#118/F03 Web SDK save → Node export independent Office byte validation',
            'generated_at':datetime.now(timezone.utc).isoformat(),'python':sys.version.split()[0],
            'libraries':{name:version(name) for name in ['openpyxl','python-docx','python-pptx','lxml']},
            'read_unit_json':False,'sdk_lo_gui_rerun':False,'scope':'existing simple fixture Office bytes; only A1/A3, DOCX prefix, PPTX first text run may change',
            'runs':runs,'total_ms':round((time.perf_counter()-started)*1000,3)}
    (HERE/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str))
    lines=['# Web SDK 保存候选：独立原格式验证','',
           '只读取现有原件和 chromium/webkit 六个原格式候选，使用 openpyxl / python-docx / python-pptx 及 ZIP/XML；未读取 unit.json，未运行 SDK、LibreOffice 或 GUI。只在 browser/ 写入脚本及这两份结果。','',
           '## 结果','', '| 引擎 | 格式 | 状态 | 通过 | 失败 | 限制/已知限制 |', '| --- | --- | --- | ---: | ---: | ---: |']
    for r in runs:
        n=r.get('counts',{})
        lines.append(f"| {r['engine']} | {r['format']} | {r['status']} | {n.get('passed',0)} | {n.get('failed',0)} | {n.get('limited',0)+n.get('known_limit',0)} |")
    lines+=['','## 断言范围','',
            '- XLSX：明细!A1=浏览器编辑、A3=键盘中文编辑；其他值/公式原样；所有公式缓存含 B5=30、B7=20、说明!B3=30；sheet 顺序/可见性、有效样式、合并、freeze、行列尺寸、对象及原格式 XML。字体 family、solid 背景色、空边框补全记录严格差异，并另查实际字体/前景/边线语义。',
            '- DOCX：首段仅前缀浏览器编辑：；所有段落与表格全文、正文 block 顺序、表格尺寸、图片字节与尺寸、页参、段落样式、页眉页脚、分页、链接。',
            '- PPTX：第 0 页第 0 元素首 run 仅替换浏览器编辑；其余全文、页序、shape 数目/名称/类型、字号/字体属性、位置/大小/旋转、图片字节/裁剪、备注、页面与核心元数据。母版/layout 与已存在 Node 无改输出对比，保持已知损失诊断；如浏览器造成进一步变化仍 failed。',
            '- 核心非易变元数据与作者 Company/Manager 等明确核对；创建/修改时间、revision、lastModifiedBy 原值及变化另记 limited。所有 ZIP parts 删除/增加/改变均记录，不能将非空 ZIP 视为成功。','', '## 实际失败与差异','']
    for r in runs:
        lines.append(f"### {r['engine']} / {r['format']}")
        lines.append('')
        if r.get('error'): lines.append('执行失败：'+json.dumps(r['error'],ensure_ascii=False))
        if r['status']=='pending': lines.append('候选未生成；生成后重跑。')
        for f in r.get('failures',[]):
            def compact(v):
                text=json.dumps(v,ensure_ascii=False,default=str)
                return text if len(text)<700 else text[:700]+'…（完整字段见 JSON）'
            provenance=f.get('provenance','not_classified')
            semantic=f.get('independent_semantic_assessment')
            lines.append(f"- **failed（原比较保留）** `{f['name']}`：期望 `{compact(f['expected'])}`；实际 `{compact(f['actual'])}`。归因：`{provenance}`"+(f"；独立语义：`{semantic}`。" if semantic else '。'))
        changes=[x for x in r.get('checks',[]) if x['status'] in ['limited','known_limit']]
        if changes: lines.append('- 限制/严格差异：'+', '.join('`'+x['name']+'`' for x in changes)+'；字段明细见 verification.json。')
        if not r.get('failures') and not r.get('error') and r['status']!='pending': lines.append('- 未发现本次语义断言失败。')
        lines.append('')
    lines+=['## Web save 链新增损失与已知转换变化','',
            '- 六个原格式文件均含预期的实际编辑；XLSX A1/A3 与三处已知公式缓存正确，DOCX 前缀及表格正文保持，PPTX 首 run 与其余页文字/字号/图片/几何保持。整体状态仍须同时考虑保留失败。',
            ('- 当前两引擎 DOCX：图片字节与尺寸断言均 passed，未发现相对 traditional Node 无编辑输出新增的保留失败。core-only 历史结果另存 verification.core-only.json/.md，其中原未编辑图片丢失的失败保持不变。'
             if all(all(v['status']=='passed' for v in r.get('checks',[]) if v['name'] in ['image_extents','images_bytes']) and not r.get('additional_failed_fields_vs_node_noop') for r in runs if r['format']=='docx' and r['status']!='pending')
             and all(r['status']!='pending' for r in runs if r['format']=='docx') else
             '- DOCX 图片与新增失败以本轮逐项结果为准；如 images_bytes/image_extents failed 且 Node 无编辑输出保留，则定位为 Web SDK save→Node export 链新增损失。原格式验证不能单独区分 save 未保留资源或 exporter 未消费资源。'),
            '- 两引擎 XLSX：A3 编辑后字体从 Calibri 11 / scheme minor / theme 色变为无显式名称字号 / FF1B1C1F；与 Node 无编辑输出不同。标记实际样式变化，不能仅按同值或等价 XML 当保留通过。',
            '- DOCX gutter=0 变为缺省、省略 PPTX 段落对齐变为显式 left，属于已核实缺省属性序列化差异；原严格比较继续 failed，追加语义说明，不称为内容丢失。',
            '- 作者/Company 等核心元数据差异以及 PPTX defaultTextStyle 丢失/尺寸类型标签变化逐项与已有 Node 输出比较；同基线者标为既有转换差异，保留原 failed，不声称是新的 Web save 失败。PPTX 母版/layout 损失独立 known_limit。',
            '', '## 边界与复跑','',
            '本脚本可证明输出文件中受测编辑和对象存在/保留的范围。浏览器实际点击或键盘操作证据由原探针负责，本脚本不推断未覆盖格式、视觉像素或复杂模板。PPTX 已知母版损失不列为新的 Web save 失败。若某对象消失或字段实际变更，结果保留 failed；失败是否源于 Web save 或已有 Node 转换需结合既有无编辑基线。',
            '', '```bash','python3 browser/check-output.py','```','',
            f"Python {report['python']}；库版本 {json.dumps(report['libraries'],ensure_ascii=False)}；本次 {report['total_ms']} ms。SHA-256、逐项断言及 ZIP 差异均在 verification.json。"]
    (HERE/'verification.md').write_text('\n'.join(lines)+'\n')
    for r in runs: print(json.dumps({k:r.get(k) for k in ['engine','format','status','counts','error']},ensure_ascii=False))
